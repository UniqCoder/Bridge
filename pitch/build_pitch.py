from pathlib import Path
from html import escape
import json
from reportlab.pdfgen import canvas
from reportlab.lib.colors import HexColor
from reportlab.lib.styles import ParagraphStyle
from reportlab.platypus import Paragraph
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
import fitz

OUT = Path(__file__).resolve().parent
ROOT = OUT.parent
D = json.loads((ROOT/'web/src/data/fixtures/detection.json').read_text())['detection']
H = json.loads((ROOT/'web/src/data/fixtures/history.json').read_text())
E = json.loads((ROOT/'web/src/data/fixtures/evidence.json').read_text())
M = json.loads((ROOT/'web/src/data/fixtures/model.json').read_text())['model']
assert D['nomad']['rule_flags'] == 382 and D['ronin']['rule_flags'] == 2
assert D['nomad']['first_volume_flag_offset_s'] == 2986
assert len(E['bundles']) == 12 and E['live']['attack_execution_blocked']

slides = []
def add(section, title, subtitle, cards, takeaway, source, notes):
    slides.append(dict(section=section,title=title,subtitle=subtitle,cards=cards,takeaway=takeaway,source=source,notes=notes))

add('01 / THE PITCH', 'BridgeWitness', 'Real-time detection of cross-chain bridge exploits | CSB-01', [
('THE PRINCIPLE','Signatures can be stolen. Conservation cannot: every destination release must consume a verified, unspent source deposit.'),
('THE PRODUCT','An independent bridge monitor that connects exact backing checks, explainable alerts and evidence an operator can re-run.'),
('THE DEMO','Public Nomad + Ronin attack replays, a seeded adversarial simulation, and evidence from a controlled two-chain gate demonstration.'),
('THE AMBITION','Move from discovering theft after settlement to holding unsupported releases before execution on integrated bridges.')],
'A valid signature is permission. A verified deposit is backing.', 'Basis: project PRD; generated history, detection and evidence fixtures.',
'Open with a question: If five legitimate validator keys approve a withdrawal, does that prove the money was ever deposited? No. BridgeWitness checks the economic fact behind the permission. Our current dashboard is a reproducible replay prototype; the local gate evidence comes from the reference backend. We are showing what works today and precisely what remains before public-network enforcement.')

add('02 / PROBLEM STATEMENT', 'Bridges trust permission. Attackers exploit the gap.', 'The critical failure: an approved withdrawal can still be economically unbacked.', [
('THE FAILURE MODE','A destination contract accepts a message or signer quorum while the matching source deposit is missing, already spent, or inconsistent.'),
('THE HUMAN LIMIT','Operators must reconcile two chains, uncertain finality and many recipients while an exploit continues. Manual review starts too late for rapid settlement.'),
('THE SCALE','On 2 August 2022, Chainalysis estimated $2B stolen across 13 bridge hacks; bridges accounted for 69% of crypto stolen that year to date.'),
('THE REQUIRED RESPONSE','Ingest live or replayed transactions; identify pair mismatches, volume spikes and signer anomalies; deliver a risk-ranked suspected exploit type.')],
'Problem statement: detect unsupported cross-chain value movement early, explain it, and enable a proportionate response.',
'External: Chainalysis, 2 Aug 2022 [S1]. Brief: CSB-01 in BridgeWitness-PRD.md.',
'A bridge is a vault spanning two chains. The problem is not simply whether a transaction looks unusual; it is whether assets leaving one side are backed by assets entering the other. Date the historical industry statistic explicitly. Do not present the 69% figure as a current market statistic. Our focus is the supported lock/mint and burn/release route model, not every possible bridge architecture.')

add('03 / WHY THE GAP MATTERS', 'Two theft patterns. One broken invariant.', 'Our replay explains why signatures and volume thresholds cannot be the final authority.', [
('RONIN: LOW EVENT COUNT','Two attack releases, only 115 seconds apart. Our configured velocity detector flags 0/2; the backing rule flags both.'),
('NOMAD: COPYCAT FAN-OUT','382 attack releases across 167 transactions and 279 recipients. The first velocity flag arrives at +2,986 seconds: 49 min 46 sec.'),
('THE COST OF WAITING','Before that Nomad volume flag, 17 releases total about $23.59M in historical dataset value. This is observed replay exposure, not money saved.'),
('OUR DESIGN RESPONSE','Use exact authorization and consumption rules for exploit verdicts. Use behavioral signals to rank and explain the developing incident.')],
'The decisive question is whether the release is backed, even when its signatures look legitimate.',
'Local: detection.json; history.json; PRD measured-results table. Values cover the dataset subset.',
'A volume detector can miss a very large theft if it consists of only two transactions. It can also detect a copycat swarm only after significant outflow. These are results from our configured baseline on these replay feeds, not a universal claim about all security products. The Nomad dataset subset totals $159.58M; it is not the complete widely reported incident valuation.')

add('04 / PROPOSED SOLUTION', 'Reconcile first. Explain every decision.', 'Three required signals, one transparent decision flow.', [
('1. CONSERVATION LEDGER','Match route + message ID; validate source contract, amount, recipient and remaining authorization. Reject replay and over-consumption with exact integer arithmetic.'),
('2. VELOCITY + SIGNER FORENSICS','Compare recent withdrawal activity with route history. Check signer membership, quorum and duplication; flag suspicious repeated signer sets when metadata exists.'),
('3. ACTIONABLE ALERTS','Name the suspected exploit, show the violated rule, expose the 0-100 priority breakdown, and correlate related releases into an incident with USD exposure.'),
('4. EVIDENCE + RESPONSE','Package the decision for browser recomputation. In the integrated gate design, a rule violation can hold a release; behavioral warnings alone cannot.')],
'Signals advise. Deterministic rules decide.',
'Local: detect/analyze.ts, classify.ts, velocity.ts; verify/accounting.ts; data/monitor.ts.',
'This directly meets the hackathon brief before introducing advanced features. The three panels are mismatched pairs, volume spikes and signer behavior. Every alert explains its exploit label and policy priority. The current monitor displays generated decisions on a virtual clock; it is not a live mainnet ingestion service. Signer checks are available for simulated/local data; the historical public CSVs contain no signer metadata.')

add('05 / THE CORE MECHANISM', 'No release without real, unspent backing.', 'The accounting rule is exact; observability determines whether we can apply it.', [
('MATCH THE AUTHORIZATION','Pin source/destination chain, bridge contract, token and message. Confirm the intended recipient and exact amount under the configured route semantics.'),
('TRACK CONSUMPTION','A deposit creates a bounded authorization. Earlier releases consume it; reusing the message cannot create a second entitlement to the same funds.'),
('CHECK THE BOUND','Prior consumption + requested release must stay within verified authorization. Token decimals and fixed-fee rules are handled in exact common units.'),
('NAME THE BLIND SPOT','If source coverage is incomplete, behind or unhealthy, say unknown. Absence from an incomplete event feed does not prove absence on-chain.')],
'Same root check, multiple attacks: fake deposits, unbacked releases, replay, amount tampering and recipient substitution.',
'Local: verify/accounting.ts; PRD backing-rule and watermark definitions. Invariant prior art: [S2].',
'Explain using a simple example: a verified deposit of 100 units cannot back a release of 101, two releases of 100, or a release to a substituted recipient. But a delayed observer must not mistake an unseen 100-unit deposit for an attack. This is why exact arithmetic and explicit coverage belong together. We credit prior research for cross-chain matching; our pitch is the combined operator workflow, evidence and response design.')

add('06 / DIFFERENTIATION', 'What makes the prototype worth remembering.', 'A defensible combination of evidence, operational judgment and reproducible evaluation.', [
('EVIDENCE YOU CAN CHALLENGE','The browser rechecks canonical JSON hash, supplied block headers, transaction/receipt inclusion and accounting. Twelve reference bundles reproduce the expected results.'),
('HONEST UNCERTAINTY','Unknown is a first-class verdict. Lag or missing coverage is visible instead of being disguised as an exploit or a safety guarantee.'),
('CONTAINMENT WITH LIMITED DISRUPTION','A route hold is demonstrated on local chains. Message-to-route-to-verifier escalation is the proposed response model, with broader gate behavior still pending.'),
('THE TRADEOFF IS VISIBLE','Counterfactual replay compares loss and disruption under five response policies. The narrow policy can lose more; the interface makes that visible.')],
'The differentiator is an auditable decision-to-response workflow, not a claim that cross-chain matching is new.',
'Local: verifier, evidence fixtures, counterfactual.ts; Architecture.tsx status table; research [S2].',
'Do not claim that no competitor has these features; no market-wide competitive audit was performed. The useful distinction is what a judge can test: alter evidence, introduce lag, compare response scopes, and inspect the reason a transfer is flagged. Proofs currently verify under supplied headers. A forged or untrusted header is not made canonical simply because its hash is internally consistent.')

add('07 / MEASURED EVIDENCE', 'Strong results, with their boundaries attached.', 'Historical replay + deterministic simulation + reference evidence artifacts.', [
('382 / 382 NOMAD RELEASES','Backing violations identified from the first offending release in the replay. 0/193 normal releases flagged in the selected 1-2 August window.'),
('2 / 2 RONIN RELEASES','Backing rule flags both historical attack releases. The configured replay volume detector flags neither.'),
('6 ATTACKS + 2 CONTROLS','All six injected attack scenarios receive the expected exploit type. A backed whale and honest traffic burst produce no exploit verdict.'),
('12 / 12 EVIDENCE BUNDLES','Generated fixtures record browser/Python agreement. Reference engine latency: p50 0.82 ms; p99 2.60 ms across 300 local pairs.')],
'Replay matching validates implementation against these labels; it does not establish independent or universal detection accuracy.',
'Local: detection.json, history.json, evidence.json, model.json; fixture-builder assertions.',
'The strongest result is not a percentage pulled from a slide: each headline links to generated artifacts and assertions. State the main caveat proactively: XChainWatcher labels come from matching rules, so agreeing with them is implementation validation. The synthetic scenarios test controlled behaviors, not future attack prevalence. Engine timing includes local SQLite commit and ML, but excludes RPC, HTTP, block time and finality. Those dominate operational detection delay.')

add('08 / THE LIVE PITCH DEMO', 'Give the judges three things to challenge.', 'A 90-second sequence using the existing console and prepared reference exhibits.', [
('00-25s: STOLEN-KEY DRAIN','Open Monitor on the seeded simulation. Seek to key compromise + drain: show valid quorum, missing backing, signer behavior and the explicit priority breakdown.'),
('25-45s: THE HONEST BURST','Show nine backed transfers causing a volume advisory. No exploit verdict: unusual activity alone is insufficient to justify containment.'),
('45-70s: TAMPER WITH EVIDENCE','Open a supplied Evidence bundle; run browser verification. Change an amount and compare with the expected digest/accounting result.'),
('70-90s: HISTORICAL + GATE EXHIBITS','Switch to the Nomad replay and show the incident exposure. Present the reference local-chain report: unbacked execution blocked; healthy settlement completed.')],
'Suggested closing challenge: can the judges make us mistake a blind spot or an honest burst for an exploit?',
'Routes: /console/monitor, /console/evidence/:sha, /console/replay. Local gate exhibit: evidence.json.',
'Rehearse the seeks and choose a bundle containing execution proofs before presenting. The current workspace contains the frontend; do not promise that a judge can trigger a fresh on-chain attack through a wired live backend. Present local gate behavior as a reference report unless that backend is running and tested. Keep the historical replay as the fallback if browser animations consume time. These timings are a proposed presentation sequence, not measured demo timings.')

add('09 / IMPACT & BENEFITS', 'Protect value without stopping every user.', 'Benefits map to specific beneficiaries and measurable operational outcomes.', [
('BRIDGE OPERATORS','Get an exact failure reason and incident-level exposure instead of manually reconciling hundreds of transaction alerts. Evidence supports faster triage.'),
('USERS + LIQUIDITY PROVIDERS','On integrated bridges, holding unsupported releases could limit reserve depletion. Healthy routes can continue where the selected scope permits.'),
('SECURITY TEAMS + AUDITORS','Receive portable evidence they can re-run and compare against the original digest; inspect which facts and trust assumptions produced the verdict.'),
('SUCCESS METRICS','Time to first correct verdict; loss before containment; legitimate transfers delayed; exploit versus advisory false alerts; evidence reproduction rate.')],
'Measure both loss avoided in modeled scenarios and the cost imposed on legitimate transfers.',
'Local: incident monitor, counterfactual UI, verifier and PRD evaluation plan. Benefits are proposed outcomes.',
'Do not say we saved $159M or $565M: we replayed historical events after the fact. Actual protection requires gate integration before execution, trusted observation, a delay window and available responders. The impact argument is specific: smaller incident workload, inspectable decisions, and the ability to choose how broadly to interrupt traffic. Future pilots should measure these separately rather than compressing them into one impressive but misleading accuracy score.')

add('10 / TECHNOLOGY STACK', 'A practical stack for exact decisions.', 'Existing frontend in this workspace; backend and local gate demonstrated in the reference project.', [
('OPERATOR INTERFACE','React 18 + TypeScript + Vite; React Router; Three.js / React Three Fiber for cross-chain flow inspection; Framer Motion + Lenis for presentation.'),
('DETECTION + STATE','TypeScript BigInt accounting and replay signals today. Reference backend: Python + FastAPI; SQLite WAL for persisted events and deterministic reconciliation.'),
('CHAIN + CRYPTOGRAPHY','EVM JSON-RPC and finalized-block scans; Solidity gate on two local Ganache chains. RLP, Keccak, Merkle-Patricia inclusion checks and canonical-JSON SHA-256.'),
('DATA + TRIAGE','SHA-256-pinned XChainWatcher CSVs; deterministic simulation seed 2410. Reference Isolation Forest is optional advisory triage, with no containment authority.')],
'Exact integer rules handle money. Cryptography makes supplied evidence inspectable. ML supplies bounded context.',
'Local: web/package.json, verify/*, fixtures; HANDOFF.md; BridgeWitness-ML-Reliability.md.',
'Explain why each layer exists rather than listing buzzwords. BigInt avoids floating-point errors in token amounts. SQLite keeps the prototype easy to reproduce. RLP and trie verification let a browser inspect execution evidence. The 3D view helps explain cross-chain flows but is not the detector. The frontend defaults to fixture mode; live SSE integration remains pending. No token, AI API or new consensus network is required for the core prototype.')

add('11 / TECHNICAL FEASIBILITY', 'The hard parts are separated and testable.', 'What exists, what is demonstrated elsewhere, and what remains.', [
('BUILT HERE','Replay monitoring, three signal panels, risk-ranked exploit alerts, incident preview, browser verifier, counterfactual policy explorer and generated data fixtures.'),
('REFERENCE DEMONSTRATION','Python engine + FastAPI, finality-tag scanner and local Solidity route hold. Artifacts show blocked unbacked execution and completed healthy settlement.'),
('NEXT INTEGRATION','Connect backend events to the frontend via SSE; add a public-testnet adapter; implement complete state-proof evidence and gate delay/escalation semantics.'),
('DEPLOYMENT CONDITIONS','Trusted/cross-checked headers and coverage, resilient RPC access, audited contracts and guardian policy. Account for finality and RPC delay before choosing the hold window.')],
'A replay demo is feasible now; public-network enforcement needs additional integration and independent security review.',
'Local: data/index.ts fixture mode; docs/Architecture.tsx; evidence.live; PRD capability status.',
'The small current build is feasible because it does not attempt a universal bridge adapter or replace consensus. However, prototype latency is not enough to prove prevention. End-to-end detection includes source finality, observer lag, service delivery and the gate reaction. A proposed 30-second optimistic window must be validated against those real delays; it is not automatically adequate. A read-only monitor remains useful before gate integration.')

add('12 / COMMERCIAL VIABILITY', 'Start with one operator and one route.', 'A focused adoption hypothesis, with no invented customers, revenue or pricing.', [
('FIRST CUSTOMER','An EVM bridge operator with explicit deposit/release semantics and an incident-response owner. Start with one supported lock/mint or burn/release route.'),
('LOWER-FRICTION PILOT','Run read-only beside existing infrastructure. Compare alerts with operator review, verify evidence and measure unknown coverage before granting hold authority.'),
('PROPOSED BUSINESS MODEL','Managed monitoring subscription priced by supported routes and operational coverage; integration/support for evidence and gate workflows. Pricing is unvalidated.'),
('WHY PAY / HOW TO VALIDATE','Potential value: less reconciliation work and reduced unsupported exposure. Interview operators, collect pilot workload data and confirm willingness to pay.')],
'First prove that an operator trusts the evidence and uses it; then evaluate optional enforcement.',
'Commercial plan: team proposal. No customer interviews, paid pilots or unit-economics evidence claimed.',
'Separate technical feasibility from commercial viability. Our buyer hypothesis is a bridge security or infrastructure team, not retail users buying a token. The main adoption barrier is granting an external system authority to interrupt settlement. Read-only operation allows operators to build confidence first. RPC, storage, operational response and route maintenance are cost drivers. Do not invent a total addressable market or declare profitable unit economics without pilot measurements.')

add('13 / NEXT MILESTONES', 'Ship evidence before broader authority.', 'Milestone gates, rather than unsupported calendar promises.', [
('M1: END-TO-END TESTNET','Connect the UI and backend on one public-testnet route. Demonstrate healthy transfer, missing backing, replay and observer-gap behavior with public transaction references.'),
('M2: STRONGER EVIDENCE','Implement verified deposit commitments and storage exclusion proofs; improve header/finality trust. Distinguish supplied-header consistency from canonical-chain assurance.'),
('M3: INDEPENDENT VALIDATION','Use separately adjudicated normal traffic and more incidents. Test reorgs, RPC gaps, malicious observers, baseline behavior and sustained load.'),
('M4: CONTROLLED OPERATOR PILOT','Begin read-only. Add audited gate integration only after latency, false holds, guardian availability and recovery procedures meet operator-agreed targets.')],
'Broaden chains and bridge models after one route is demonstrably correct and operationally useful.',
'Roadmap synthesized from PRD, HANDOFF and ML reliability contract; milestones not completed claims.',
'The next steps deliberately strengthen trust, not visual polish. A production gate creates availability and governance risks, so its authority should expand only after the monitor has been independently validated. Liquidity and intent bridges need their own semantics and adapters; they are outside the current exact-lock model. We do not assign arbitrary two-week deadlines to work requiring audits or external operators.')

add('14 / HONEST LIMITS', 'Trust is earned by exposing the assumptions.', 'These boundaries belong in the pitch, not hidden after a judge asks.', [
('DATASET LIMIT','Historical labels derive from matching rules; replay coverage is assumed complete. Public CSVs have no signer data. Dollar values are historical dataset values.'),
('PROOF LIMIT','Current evidence proves internal consistency and inclusion under supplied headers. State non-deposit proof integration and stronger consensus trust remain roadmap.'),
('DEPLOYMENT LIMIT','UI defaults to fixtures; gate artifacts use local, unaudited contracts. No public-chain prevention, recovered funds or production uptime is claimed.'),
('MODEL LIMIT','Reference forest warning recall is 9/384 = 2.34% on the held-out split. It advises only; it cannot override accounting or authorize containment.')],
'We can defend every demonstrated result without claiming that two incidents predict every future exploit.',
'Local: SiteLayout.tsx NON_CLAIMS; model.json; reliability contract; architecture status.',
'A strong answer acknowledges limitations before they undermine the story. ML has a high raw-score AUC in the artifacts but poor warning recall at its operating threshold, so we do not turn AUC into an accuracy headline. Watermarks from a lying observer can create false confidence; coverage claims need a trustworthy source or stronger evidence. Proof of non-deposit is not complete in this frontend. These are engineering milestones, not reasons to invent a stronger current result.')

add('15 / THE CLOSE', 'Every withdrawal should carry its witness.', 'BridgeWitness: exact backing checks, visible uncertainty and inspectable response.', [
('THE PROBLEM WE SOLVE','A bridge can approve a release that was never properly backed. Valid signatures and dramatic volume are insufficient evidence of legitimacy.'),
('THE WORKING PROTOTYPE','A required three-signal dashboard, reproducible public-attack replay, controlled attack scenarios and independently re-executable supplied evidence.'),
('THE REASON TO REMEMBER US','We show why a release fails, when the observer cannot decide, and what a narrower or broader response would cost.'),
('THE NEXT ASK','A design partner for one EVM route: read-only validation first, independently reviewed evidence next, and audited selective enforcement after that.')],
'Do not just ask who signed. Ask what backs the release.',
'Pitch positioning based on demonstrated prototype and explicitly labeled roadmap.',
'Close without adding new performance claims. We satisfy the monitoring brief today and demonstrate a path from alerts to accountable containment. The ask is a pilot and technical feedback, not an assertion of adoption. Final line: Every withdrawal should carry its witness. Leave the Evidence view ready for questions. For a shorter five-minute presentation, prioritize slides 1-7, 9-12 and 15, then use the remaining slides as backup.')

add('A1 / DETAILED PROBLEM & SOLUTION', 'The project in submission-ready language.', 'Use these paragraphs for the hackathon application or abstract.', [
('PROBLEM STATEMENT','Cross-chain bridges can release destination assets without a legitimate, unspent source authorization. Compromised signers, replayed messages and inconsistent transfers exploit this gap while operators face fragmented chain data and fast outflow.'),
('PROPOSED SOLUTION','BridgeWitness reconciles each release against exact source backing, adds velocity and signer context, and emits an explained exploit-priority alert. It exposes incomplete coverage instead of forcing a verdict and packages the supplied evidence for browser verification.'),
('EXPECTED IMPACT','Operators gain a reproducible basis for incident triage. With pre-execution gate integration, unsupported releases could be held while healthy traffic continues under a narrower response scope, reducing potential reserve loss and unnecessary disruption.'),
('FEASIBILITY + VIABILITY','A working replay UI and local reference gate evidence establish prototype feasibility. Adoption starts with read-only validation on one EVM route; public integration, independent evaluation and audited enforcement remain the next milestones.')],
'Submission angle: economic correctness plus auditable operator action.',
'Submission wording synthesized from the project; expected benefits and business adoption remain hypotheses.',
'This appendix is deliberately complete enough to paste into a submission form. Shorten by removing details, not by converting potential benefits into realized outcomes. If a form has separate impact and novelty boxes, emphasize operational metrics in impact and the combination of reproducible evidence, observability and scope tradeoffs in novelty. Credit XChainWatcher for the matching research and data.')

add('A2 / HOW AN ALERT EARNS ITS PRIORITY', 'A score you can explain in one sentence.', 'Example from the seeded key-compromise-and-drain simulation.', [
('THE OBSERVATION','A WETH-route release requests 457 WETH, valued at $1,354,457 in the simulation. Five authorized signers co-sign, but no matching source authorization backs it.'),
('THE CLASSIFICATION','Suspected validator key compromise: valid quorum on an unbacked release. The example also shows lockstep signers and a volume deviation of 12.53.'),
('THE SCORE','Base exploit priority 80 + capped exposure term 15 + campaign term 5 = 100, with evidence completeness 1.0. This is a policy ranking, not a probability.'),
('THE AUTHORITY BOUNDARY','The backing violation supplies the containment basis. Volume and signer behavior add context; a high behavioral score alone cannot authorize a hold.')],
'The alert explains what failed, what is suspected and why it comes first in the queue.',
'Exact example: detection.json -> detection.example_alert. Synthetic, seed 2410.',
'Judges often ask what the risk score actually means. Show the arithmetic rather than claiming a calibrated chance of theft. The general exposure term is min(15, 3*log10(USD)); campaign adds at most five and the result is bounded to 0-100 before evidence scaling. Suspected key compromise is an explanatory label, not forensic proof that a private key was stolen. Behavioral evidence is metadata and can be evaded.')

add('A3 / FEASIBILITY UNDER FAILURE', 'A security tool must know when to abstain.', 'Planned operational behavior and the remaining trust boundary.', [
('CONGESTION / SOURCE LAG','Delay alone yields pending, overdue or unknown according to observation state. Do not manufacture a missing-deposit exploit from an incomplete source scan.'),
('RPC GAPS / REORGS','Expose coverage health, rebuild affected accounting on canonical changes and suppress advisory scoring when evidence is unresolved. Broader end-to-end scenario evaluation is pending.'),
('LYING OBSERVER','A watermark is an observer assertion, not cryptographic proof of non-deposit. A malicious data source remains a documented limitation until stronger proofs and chain trust are integrated.'),
('GATE OUTAGE / BAD GUARDIAN','A pre-execution delay introduces availability and governance tradeoffs. Define accountable hold/resolve actions and audit them before permitting production authority.')],
'Abstention prevents false certainty; it does not by itself guarantee protection or settlement availability.',
'Local: reliability contract; PRD trust/non-claims; Eval.tsx scenarios marked pending.',
'An unknown verdict must be visible to the operator and must not be presented as safe. This slide states intended behavior without claiming the full backend adversarial suite is complete in this workspace. Depending on gate policy, waiting for evidence may itself delay legitimate users. Independent incident labels, adversarial observer tests and consensus-grade header validation are necessary to judge production readiness.')

add('A4 / JUDGE QUESTIONS', 'Answers that survive technical scrutiny.', 'Keep these responses concise, then open the supporting artifact.', [
('IS THIS JUST XCHAINWATCHER?','We reuse and credit its matching principle and data. Our prototype combines the required signals with browser evidence, visible coverage, incident triage and response-policy comparisons; integrated containment remains partly roadmap.'),
('DID YOU PREVENT NOMAD?','No. We replayed the historical dataset. A separate controlled local gate demonstration blocked an unsupported release; counterfactual losses are modeled estimates.'),
('WHY NOT USE AI AS THE DETECTOR?','The reference forest flags only 9/384 held-out attacks at its threshold. Exact accounting supplies the decision; the model is bounded advisory triage with explicit abstention.'),
('CAN IT COVER EVERY BRIDGE?','No. Current semantics cover configured lock/mint and burn/release routes with fixed fees. Liquidity and intent designs require additional adapters and correctness rules.')],
'If asked about production readiness: public feed wiring, stronger trust, independent evaluation and audited gates are still required.',
'Local: PRD; NON_CLAIMS; model.json; reliability contract; research [S2].',
'For differentiation questions, demonstrate the evidence workflow rather than dismissing existing security tools. For scalability, quote only the local 300-pair benchmark and say long-running or concurrent-writer load was not evaluated. For business questions, call the managed-monitoring model a hypothesis. For cryptographic questions, separate inclusion under a supplied header from external assurance that the header is canonical.')

add('A5 / SOURCES & CLAIM HYGIENE', 'A pitch the team can trace back to evidence.', 'External context is cited; prototype measurements come from local artifacts.', [
('[S1] HISTORICAL INDUSTRY CONTEXT','Chainalysis, 2 August 2022: Cross-Chain Bridge Hacks Emerge as Top Security Risk. The $2B / 13 hacks / 69% figures are historical context, not current statistics.'),
('[S2] RESEARCH + DATA CREDIT','Augusto et al., XChainWatcher: Monitoring and Identifying Attacks in Cross-Chain Bridges, arXiv:2410.02029. Dataset attribution is retained in the project notices.'),
('LOCAL MEASUREMENT SOURCES','web/src/data/fixtures/{history,detection,evidence,model}.json; generated by web/scripts/build-fixtures.mjs. PRD and HANDOFF describe status and assumptions.'),
('PRESENTATION RULE','Say on this replay, local benchmark, supplied evidence, proposed benefit or modeled loss. Avoid 100% accurate, saved funds, universal prevention or live-mainnet deployment.')],
'Source links and detailed presenter notes accompany the editable deck.',
'[S1] chainalysis.com/blog/cross-chain-bridge-hacks-2022/ | [S2] arxiv.org/abs/2410.02029',
'External sources were checked for historical context and research attribution. Local measurements were inspected in generated fixtures; the full application, fixture-generation pipeline and reference-chain demonstration were not rerun while preparing this deck. Do not interpret slide preparation as a security audit or validation of production protection. Presenter notes contain the qualifications needed for each headline.')

COL = dict(bg='#0B1420', panel='#152333', ink='#F2F5F7', mute='#A8B6C5', mint='#67E3BF', line='#294051')
pdfmetrics.registerFont(TTFont('Arial', 'C:/Windows/Fonts/arial.ttf'))
pdfmetrics.registerFont(TTFont('Arial-Bold', 'C:/Windows/Fonts/arialbd.ttf'))
W, HEIGHT = 960, 540
pdf_path = OUT/'BridgeWitness-Hackathon-Pitch.pdf'
c = canvas.Canvas(str(pdf_path),pagesize=(W,HEIGHT))
c.setTitle('BridgeWitness | Hackathon Pitch | CSB-01')
c.setAuthor('BridgeWitness Team')
prs = Presentation()
prs.slide_width = Inches(13.333333)
prs.slide_height = Inches(7.5)

def ptext(slide, x,y,w,h,text,size,color,bold=False):
    box=slide.shapes.add_textbox(Inches(x/72),Inches(y/72),Inches(w/72),Inches(h/72))
    tf=box.text_frame; tf.word_wrap=True
    tf.margin_left=tf.margin_right=Pt(0)
    tf.margin_top=tf.margin_bottom=Pt(0)
    para=tf.paragraphs[0]; para.text=text
    para.font.name='Arial'; para.font.size=Pt(size); para.font.bold=bold
    para.font.color.rgb=RGBColor.from_string(color.lstrip('#'))
    para.space_after=Pt(0)

def rect(slide,x,y,w,h,color):
    shape=slide.shapes.add_shape(1,Inches(x/72),Inches(y/72),Inches(w/72),Inches(h/72))
    shape.fill.solid(); shape.fill.fore_color.rgb=RGBColor.from_string(color.lstrip('#')); shape.line.fill.background()
    c.setFillColor(HexColor(color)); c.rect(x,HEIGHT-y-h,w,h,fill=1,stroke=0)

def text(slide,x,y,w,h,txt,size=15,color=None,bold=False):
    color=color or COL['ink']
    style=ParagraphStyle('body',fontName='Arial-Bold' if bold else 'Arial',fontSize=size,leading=size*1.27,textColor=HexColor(color))
    p=Paragraph(escape(txt),style); pw,ph=p.wrap(w,h)
    assert ph <= h, (txt,ph,h)
    p.drawOn(c,x,HEIGHT-y-ph)
    ptext(slide,x,y,w,h,txt,size,color,bold)

for i,s in enumerate(slides,1):
    slide=prs.slides.add_slide(prs.slide_layouts[6])
    slide.background.fill.solid(); slide.background.fill.fore_color.rgb=RGBColor.from_string(COL['bg'][1:])
    c.setFillColor(HexColor(COL['bg'])); c.rect(0,0,W,HEIGHT,fill=1,stroke=0)
    rect(slide,38,33,5,18,COL['mint'])
    text(slide,54,34,850,18,s['section'],11,COL['mint'],True)
    text(slide,38,71,885,72,s['title'],32 if len(s['title'])<49 else 29,bold=True)
    text(slide,38,145,882,39,s['subtitle'],14,COL['mute'])
    for k,(label,body) in enumerate(s['cards']):
        x=38+(k%2)*449; y=197+(k//2)*119
        rect(slide,x,y,435,106,COL['panel'])
        text(slide,x+15,y+12,406,25,label,11,COL['mint'],True)
        text(slide,x+15,y+37,406,66,body,12.5)
    text(slide,38,444,880,45,s['takeaway'],15,COL['ink'],True)
    text(slide,38,502,840,26,s['source'],8,COL['mute'])
    text(slide,902,502,40,18,f'{i:02d}',11,COL['mint'],True)
    slide.notes_slide.notes_text_frame.text=s['notes']+'\n\nEvidence: '+s['source']
    c.showPage()
c.save()
prs.save(OUT/'BridgeWitness-Hackathon-Pitch.pptx')

md=['# BridgeWitness — Hackathon pitch content and presenter notes','',
    '20 slides: 15 main slides + 5 backup slides. Suggested shorter delivery: slides 1–7, 9–12 and 15. All monetary figures are historical dataset or simulated values. Current UI is fixture-based; local gate results are reference artifacts.','',
    '## Ready-to-speak opening (about 30 seconds)','',
    'If five legitimate validators sign a bridge withdrawal, does that prove the money was ever deposited? No. Stolen keys can produce valid signatures. BridgeWitness asks a harder question: does every release consume a real, unspent deposit? We built a three-signal monitoring dashboard, replayed Nomad and Ronin, and made the evidence re-executable in the browser. Our goal is to turn an explained violation into an accountable hold before funds move on an integrated bridge.','']
for i,s in enumerate(slides,1):
    md += [f"## Slide {i}: {s['title']}",'',s['subtitle'],'']
    for label,body in s['cards']: md += [f'**{label}** — {body}','']
    md += [f"**Takeaway:** {s['takeaway']}",'',f"**Presenter notes:** {s['notes']}",'',f"**Evidence:** {s['source']}",'']
md += ['## External sources','',
       '- [S1: Chainalysis, 2 August 2022](https://www.chainalysis.com/blog/cross-chain-bridge-hacks-2022/)',
       '- [S2: XChainWatcher research paper](https://arxiv.org/abs/2410.02029)','',
       '## Final preparation','',
       'Add team names if required. Rehearse the console seeks, evidence bundle and fallback exhibits. Treat the local gate report as an exhibit unless the reference backend is running. Retain the dataset/proof caveats whenever reusing numerical claims.']
(OUT/'BridgeWitness-Presenter-Notes.md').write_text('\n'.join(md),encoding='utf-8')

# Self-check the exported artifact, including page count, text extraction and editable slides.
doc=fitz.open(pdf_path)
assert len(doc)==len(slides)==20
assert len(Presentation(OUT/'BridgeWitness-Hackathon-Pitch.pptx').slides)==20
for page in doc:
    assert 'BridgeWitness' in page.get_text() or len(page.get_text())>500
    for block in page.get_text('blocks'):
        assert block[0]>=0 and block[1]>=0 and block[2]<=W and block[3]<=HEIGHT, block
preview=fitz.open()
for n in [0,2,6,9,11,13]:
    pix=doc[n].get_pixmap(matrix=fitz.Matrix(.7,.7))
    page=preview.new_page(width=W*.7,height=HEIGHT*.7)
    page.insert_image(page.rect,stream=pix.tobytes('png'))
preview.save(OUT/'Pitch-Preview.pdf')
doc[6].get_pixmap(matrix=fitz.Matrix(1.2,1.2)).save(OUT/'slide-preview.png')
print('Created and checked: 20-page PDF, 20-slide editable PowerPoint, detailed presenter notes.')
print('Output folder:',OUT)
